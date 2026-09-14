#!/usr/bin/env python3
"""Extrai candidatos de PDFs e planilhas Excel (.xlsx); não contém regras de bloqueio nem toca no banco."""
import argparse
import gc
import json
import re
import sys
from pathlib import Path

DOMAIN = re.compile(r"(?<![\w.-])(?:[a-z0-9](?:[a-z0-9-]{0,61}[a-z0-9])?\.)+[a-z]{2,63}(?![\w.-])", re.I)
GLUED_PROSE = re.compile(
    r"^(?P<domain>.+\.(?:com|net|org|app|site|online|xyz|tv|io|me|info|biz|com\.br))"
    r"(?:todas|todos|esta|este|essas|esses|para|conforme|atraves|acima|abaixo)$",
    re.I,
)

def normalize_candidate(value):
    """Separa palavras do texto que o PDF colou depois de um domínio."""
    value = value.lower().rstrip('.')
    glued = GLUED_PROSE.fullmatch(value)
    return glued.group('domain') if glued else value

def candidates(text):
    return {normalize_candidate(match.group(0)) for match in DOMAIN.finditer(text or '')}

def extract_pdf(path):
    import pdfplumber
    found, candidate_count, invalid = set(), 0, 0
    pages = 0
    with pdfplumber.open(path) as pdf:
        pages = len(pdf.pages)
        for page in pdf.pages:
            page_found = set()
            try:
                for table in page.extract_tables() or []:
                    for row in table or []:
                        for cell in row or []:
                            page_found.update(candidates(cell or ''))
                # Fallback obrigatório: texto sempre complementa tabelas incompletas.
                page_found.update(candidates(page.extract_text(x_tolerance=2, y_tolerance=3) or ''))
                candidate_count += len(page_found)
                found.update(page_found)
            finally:
                if hasattr(page, 'flush_cache'):
                    page.flush_cache()
                for attr in ('_objects', '_layout'):
                    if hasattr(page, attr):
                        setattr(page, attr, None)
                del page_found
                gc.collect()
    return {'filename': Path(path).name, 'pages': pages, 'candidates': candidate_count, 'domains': len(found), 'invalid': invalid}, found

def extract_xlsx(path):
    import openpyxl
    found, candidate_count = set(), 0
    sheets = 0
    # read_only + data_only: streaming (nao carrega a planilha inteira na memoria) e
    # pega o valor calculado de formulas, nao a formula em si.
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = len(workbook.sheetnames)
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    value = cell.value
                    if value is None:
                        continue
                    candidate_count += 1
                    found.update(candidates(str(value)))
    finally:
        workbook.close()
    return {'filename': Path(path).name, 'pages': sheets, 'candidates': candidate_count, 'domains': len(found), 'invalid': 0}, found

def extract(path):
    suffix = Path(path).suffix.lower()
    if suffix in ('.xlsx', '.xlsm'):
        return extract_xlsx(path)
    return extract_pdf(path)

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('pdfs', nargs='+')
    args = parser.parse_args()
    files, domains = [], set()
    try:
        for filename in args.pdfs:
            item, extracted = extract(filename)
            files.append(item); domains.update(extracted)
        print(json.dumps({'status':'ok','files':files,'domains':sorted(domains)}, ensure_ascii=True))
        return 0
    except Exception as exc:
        print(json.dumps({'status':'error','error':type(exc).__name__}), file=sys.stdout)
        return 1
if __name__ == '__main__':
    raise SystemExit(main())
